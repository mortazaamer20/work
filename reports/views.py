import csv
import io
from datetime import date, timedelta
from decimal import Decimal

from django.shortcuts import render, get_object_or_404
from django.http import HttpResponse
from django.db.models import Avg, Count, Q, Min, Max, F

from accounts.views import require_perm
from beneficiaries.models import Beneficiary
from assessments.models import DailyAssessment, RehabilitationIndex, ActivityRecord, ClinicRecord
from centers.models import Center, Activity, Clinic

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False


@require_perm("reports", "view")
def reports_home(request):
    centers = Center.objects.filter(is_active=True)
    beneficiaries = Beneficiary.objects.filter(status="active")
    return render(request, "reports/home.html", {
        "centers": centers,
        "beneficiaries": beneficiaries,
    })


@require_perm("reports", "view")
def individual_report(request):
    beneficiary_id = request.GET.get("beneficiary")
    date_from = request.GET.get("date_from", "")
    date_to = request.GET.get("date_to", "")
    export_format = request.GET.get("export", "")

    if not beneficiary_id:
        return render(request, "reports/individual_select.html", {
            "beneficiaries": Beneficiary.objects.filter(status="active"),
        })

    b = get_object_or_404(Beneficiary, pk=beneficiary_id)
    qs = RehabilitationIndex.objects.filter(
        assessment__beneficiary=b
    ).select_related("assessment").order_by("-assessment__date")

    if date_from:
        qs = qs.filter(assessment__date__gte=date_from)
    if date_to:
        qs = qs.filter(assessment__date__lte=date_to)

    indices = list(qs)

    stats = _calculate_individual_stats(indices)

    activity_records = ActivityRecord.objects.filter(
        assessment__beneficiary=b
    ).select_related("activity__center")
    if date_from:
        activity_records = activity_records.filter(assessment__date__gte=date_from)
    if date_to:
        activity_records = activity_records.filter(assessment__date__lte=date_to)

    activity_stats = activity_records.values(
        "activity__name", "activity__center__name"
    ).annotate(
        avg_score=Avg("score"),
        count=Count("id"),
    ).order_by("-avg_score")

    if export_format == "excel" and HAS_OPENPYXL:
        return _export_individual_excel(b, indices, stats, activity_stats, date_from, date_to)
    elif export_format == "csv":
        return _export_individual_csv(b, indices, date_from, date_to)

    return render(request, "reports/individual_report.html", {
        "b": b, "indices": indices, "stats": stats,
        "activity_stats": activity_stats,
        "date_from": date_from, "date_to": date_to,
    })


@require_perm("reports", "view")
def center_report(request):
    center_id = request.GET.get("center")
    date_from = request.GET.get("date_from", "")
    date_to = request.GET.get("date_to", "")
    export_format = request.GET.get("export", "")

    if not center_id:
        return render(request, "reports/center_select.html", {
            "centers": Center.objects.filter(is_active=True),
        })

    center = get_object_or_404(Center, pk=center_id)

    activity_records = ActivityRecord.objects.filter(
        activity__center=center
    ).select_related("activity", "assessment__beneficiary")
    clinic_records = ClinicRecord.objects.filter(
        clinic__center=center
    ).select_related("clinic", "assessment__beneficiary")

    if date_from:
        activity_records = activity_records.filter(assessment__date__gte=date_from)
        clinic_records = clinic_records.filter(assessment__date__gte=date_from)
    if date_to:
        activity_records = activity_records.filter(assessment__date__lte=date_to)
        clinic_records = clinic_records.filter(assessment__date__lte=date_to)

    activity_stats = activity_records.values("activity__name").annotate(
        avg_score=Avg("score"), count=Count("id"),
        unique_beneficiaries=Count("assessment__beneficiary", distinct=True),
    ).order_by("activity__name")

    clinic_stats = clinic_records.values("clinic__name").annotate(
        avg_score=Avg("score"), count=Count("id"),
        unique_beneficiaries=Count("assessment__beneficiary", distinct=True),
    ).order_by("clinic__name")

    beneficiary_performance = activity_records.values(
        "assessment__beneficiary__full_name", "assessment__beneficiary__ref_number"
    ).annotate(
        avg_score=Avg("score"), total_sessions=Count("id"),
    ).order_by("-avg_score")[:20]

    daily_trend = activity_records.values("assessment__date").annotate(
        avg_score=Avg("score"), count=Count("id"),
    ).order_by("assessment__date")

    if export_format == "excel" and HAS_OPENPYXL:
        return _export_center_excel(center, activity_stats, clinic_stats, beneficiary_performance, daily_trend, date_from, date_to)
    elif export_format == "csv":
        return _export_center_csv(center, activity_stats, clinic_stats, date_from, date_to)

    return render(request, "reports/center_report.html", {
        "center": center, "activity_stats": activity_stats, "clinic_stats": clinic_stats,
        "beneficiary_performance": beneficiary_performance, "daily_trend": list(daily_trend),
        "date_from": date_from, "date_to": date_to,
    })


@require_perm("reports", "view")
def group_report(request):
    date_from = request.GET.get("date_from", "")
    date_to = request.GET.get("date_to", "")
    status_filter = request.GET.get("status", "active")
    export_format = request.GET.get("export", "")

    qs = RehabilitationIndex.objects.select_related("assessment__beneficiary")
    if date_from:
        qs = qs.filter(assessment__date__gte=date_from)
    if date_to:
        qs = qs.filter(assessment__date__lte=date_to)
    if status_filter:
        qs = qs.filter(assessment__beneficiary__status=status_filter)

    overall_stats = qs.aggregate(
        avg_total=Avg("total_score"),
        avg_medical=Avg("medical_psychological_score"),
        avg_behavioral=Avg("behavioral_score"),
        avg_skill=Avg("skill_technical_score"),
        avg_social=Avg("social_spiritual_score"),
        total_assessments=Count("id"),
    )

    status_distribution = qs.values("status").annotate(count=Count("id")).order_by("status")

    per_beneficiary = qs.values(
        "assessment__beneficiary__full_name",
        "assessment__beneficiary__ref_number",
        "assessment__beneficiary__pk",
    ).annotate(
        avg_score=Avg("total_score"),
        sessions=Count("id"),
        latest_status=Max("status"),
        avg_medical=Avg("medical_psychological_score"),
        avg_behavioral=Avg("behavioral_score"),
        avg_skill=Avg("skill_technical_score"),
        avg_social=Avg("social_spiritual_score"),
    ).order_by("-avg_score")

    daily_trend = qs.values("assessment__date").annotate(
        avg_score=Avg("total_score"),
        count=Count("id"),
    ).order_by("assessment__date")

    if export_format == "excel" and HAS_OPENPYXL:
        return _export_group_excel(overall_stats, status_distribution, per_beneficiary, daily_trend, date_from, date_to)
    elif export_format == "csv":
        return _export_group_csv(per_beneficiary, date_from, date_to)

    return render(request, "reports/group_report.html", {
        "overall_stats": overall_stats, "status_distribution": status_distribution,
        "per_beneficiary": per_beneficiary, "daily_trend": list(daily_trend),
        "date_from": date_from, "date_to": date_to, "status_filter": status_filter,
    })


@require_perm("reports", "view")
def period_report(request):
    date_from = request.GET.get("date_from", "")
    date_to = request.GET.get("date_to", "")
    export_format = request.GET.get("export", "")

    if not date_from or not date_to:
        today = date.today()
        return render(request, "reports/period_select.html", {
            "default_from": (today - timedelta(days=30)).isoformat(),
            "default_to": today.isoformat(),
        })

    qs = RehabilitationIndex.objects.filter(
        assessment__date__gte=date_from, assessment__date__lte=date_to
    ).select_related("assessment__beneficiary")

    daily_stats = qs.values("assessment__date").annotate(
        avg_score=Avg("total_score"),
        count=Count("id"),
        progressing=Count("id", filter=Q(status="progressing")),
        stable=Count("id", filter=Q(status="stable")),
        regressing=Count("id", filter=Q(status="regressing")),
    ).order_by("assessment__date")

    center_performance = []
    for center in Center.objects.filter(is_active=True):
        a_avg = ActivityRecord.objects.filter(
            activity__center=center,
            assessment__date__gte=date_from, assessment__date__lte=date_to
        ).aggregate(avg=Avg("score"))["avg"]
        c_avg = ClinicRecord.objects.filter(
            clinic__center=center,
            assessment__date__gte=date_from, assessment__date__lte=date_to
        ).aggregate(avg=Avg("score"))["avg"]
        avg_score = a_avg or c_avg or 0
        center_performance.append({
            "name": center.name, "axis": center.get_axis_display(),
            "avg_score": round(avg_score, 1) if avg_score else 0,
        })

    overall = qs.aggregate(
        avg_total=Avg("total_score"),
        avg_medical=Avg("medical_psychological_score"),
        avg_behavioral=Avg("behavioral_score"),
        avg_skill=Avg("skill_technical_score"),
        avg_social=Avg("social_spiritual_score"),
        total=Count("id"),
    )

    if export_format == "excel" and HAS_OPENPYXL:
        return _export_period_excel(daily_stats, center_performance, overall, date_from, date_to)

    return render(request, "reports/period_report.html", {
        "daily_stats": list(daily_stats), "center_performance": center_performance,
        "overall": overall, "date_from": date_from, "date_to": date_to,
    })


# === Helper functions ===

def _calculate_individual_stats(indices):
    if not indices:
        return {"count": 0}
    scores = [float(i.total_score) for i in indices]
    return {
        "count": len(scores),
        "avg": round(sum(scores) / len(scores), 1),
        "min": round(min(scores), 1),
        "max": round(max(scores), 1),
        "latest": round(scores[0], 1) if scores else 0,
        "trend": "up" if len(scores) >= 2 and scores[0] > scores[1] else ("down" if len(scores) >= 2 and scores[0] < scores[1] else "stable"),
        "avg_medical": round(sum(float(i.medical_psychological_score) for i in indices) / len(indices), 1),
        "avg_behavioral": round(sum(float(i.behavioral_score) for i in indices) / len(indices), 1),
        "avg_skill": round(sum(float(i.skill_technical_score) for i in indices) / len(indices), 1),
        "avg_social": round(sum(float(i.social_spiritual_score) for i in indices) / len(indices), 1),
    }


def _make_excel_response(wb, filename):
    response = HttpResponse(content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    wb.save(response)
    return response


def _style_header(ws, row=1):
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="0891B2", end_color="0891B2", fill_type="solid")
    for cell in ws[row]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")


def _export_individual_excel(b, indices, stats, activity_stats, date_from, date_to):
    wb = Workbook()

    ws = wb.active
    ws.title = "مؤشر التأهيل"
    ws.sheet_view.rightToLeft = True
    headers = ["التاريخ", "الطبي والنفسي", "السلوكي", "المهاري", "الاجتماعي", "المؤشر الكلي", "الحالة"]
    ws.append(headers)
    _style_header(ws)
    for idx in indices:
        ws.append([
            str(idx.assessment.date),
            float(idx.medical_psychological_score),
            float(idx.behavioral_score),
            float(idx.skill_technical_score),
            float(idx.social_spiritual_score),
            float(idx.total_score),
            idx.get_status_display(),
        ])

    ws2 = wb.create_sheet("أداء الأنشطة")
    ws2.sheet_view.rightToLeft = True
    ws2.append(["المركز", "النشاط", "متوسط الدرجة", "عدد الجلسات"])
    _style_header(ws2)
    for s in activity_stats:
        ws2.append([
            s["activity__center__name"],
            s["activity__name"],
            round(s["avg_score"], 1) if s["avg_score"] else 0,
            s["count"],
        ])

    ws3 = wb.create_sheet("ملخص")
    ws3.sheet_view.rightToLeft = True
    ws3.append(["المستفيد", b.full_name])
    ws3.append(["الرقم المرجعي", b.ref_number])
    ws3.append(["الفترة", f"{date_from or 'الكل'} إلى {date_to or 'الكل'}"])
    ws3.append(["عدد التقييمات", stats.get("count", 0)])
    ws3.append(["المتوسط العام", stats.get("avg", 0)])
    ws3.append(["أعلى مؤشر", stats.get("max", 0)])
    ws3.append(["أدنى مؤشر", stats.get("min", 0)])

    return _make_excel_response(wb, f"تقرير_{b.ref_number}_{date.today()}.xlsx")


def _export_individual_csv(b, indices, date_from, date_to):
    response = HttpResponse(content_type="text/csv; charset=utf-8-sig")
    response["Content-Disposition"] = f'attachment; filename="report_{b.ref_number}_{date.today()}.csv"'
    response.write("﻿")
    writer = csv.writer(response)
    writer.writerow(["التاريخ", "الطبي والنفسي", "السلوكي", "المهاري", "الاجتماعي", "المؤشر الكلي", "الحالة"])
    for idx in indices:
        writer.writerow([
            str(idx.assessment.date), float(idx.medical_psychological_score),
            float(idx.behavioral_score), float(idx.skill_technical_score),
            float(idx.social_spiritual_score), float(idx.total_score), idx.get_status_display(),
        ])
    return response


def _export_center_excel(center, activity_stats, clinic_stats, beneficiary_performance, daily_trend, date_from, date_to):
    wb = Workbook()
    ws = wb.active
    ws.title = "أداء الأنشطة"
    ws.sheet_view.rightToLeft = True
    ws.append(["النشاط", "متوسط الدرجة", "عدد الجلسات", "عدد المستفيدين"])
    _style_header(ws)
    for s in activity_stats:
        ws.append([s["activity__name"], round(s["avg_score"] or 0, 1), s["count"], s["unique_beneficiaries"]])

    if clinic_stats:
        ws2 = wb.create_sheet("أداء العيادات")
        ws2.sheet_view.rightToLeft = True
        ws2.append(["العيادة", "متوسط الدرجة", "عدد الجلسات", "عدد المستفيدين"])
        _style_header(ws2)
        for s in clinic_stats:
            ws2.append([s["clinic__name"], round(s["avg_score"] or 0, 1), s["count"], s["unique_beneficiaries"]])

    ws3 = wb.create_sheet("أداء المستفيدين")
    ws3.sheet_view.rightToLeft = True
    ws3.append(["الاسم", "الرقم المرجعي", "متوسط الدرجة", "عدد الجلسات"])
    _style_header(ws3)
    for p in beneficiary_performance:
        ws3.append([
            p["assessment__beneficiary__full_name"],
            p["assessment__beneficiary__ref_number"],
            round(p["avg_score"] or 0, 1), p["total_sessions"],
        ])

    return _make_excel_response(wb, f"تقرير_{center.slug}_{date.today()}.xlsx")


def _export_center_csv(center, activity_stats, clinic_stats, date_from, date_to):
    response = HttpResponse(content_type="text/csv; charset=utf-8-sig")
    response["Content-Disposition"] = f'attachment; filename="report_{center.slug}_{date.today()}.csv"'
    response.write("﻿")
    writer = csv.writer(response)
    writer.writerow(["النوع", "الاسم", "متوسط الدرجة", "عدد الجلسات", "عدد المستفيدين"])
    for s in activity_stats:
        writer.writerow(["نشاط", s["activity__name"], round(s["avg_score"] or 0, 1), s["count"], s["unique_beneficiaries"]])
    for s in clinic_stats:
        writer.writerow(["عيادة", s["clinic__name"], round(s["avg_score"] or 0, 1), s["count"], s["unique_beneficiaries"]])
    return response


def _export_group_excel(overall_stats, status_distribution, per_beneficiary, daily_trend, date_from, date_to):
    wb = Workbook()
    ws = wb.active
    ws.title = "أداء المستفيدين"
    ws.sheet_view.rightToLeft = True
    ws.append(["الاسم", "الرقم المرجعي", "المتوسط", "عدد التقييمات", "الطبي", "السلوكي", "المهاري", "الاجتماعي"])
    _style_header(ws)
    for p in per_beneficiary:
        ws.append([
            p["assessment__beneficiary__full_name"], p["assessment__beneficiary__ref_number"],
            round(p["avg_score"] or 0, 1), p["sessions"],
            round(p["avg_medical"] or 0, 1), round(p["avg_behavioral"] or 0, 1),
            round(p["avg_skill"] or 0, 1), round(p["avg_social"] or 0, 1),
        ])

    ws2 = wb.create_sheet("ملخص عام")
    ws2.sheet_view.rightToLeft = True
    ws2.append(["الفترة", f"{date_from or 'الكل'} إلى {date_to or 'الكل'}"])
    ws2.append(["إجمالي التقييمات", overall_stats.get("total_assessments", 0)])
    ws2.append(["المتوسط العام", round(overall_stats.get("avg_total") or 0, 1)])
    ws2.append(["متوسط الطبي والنفسي", round(overall_stats.get("avg_medical") or 0, 1)])
    ws2.append(["متوسط السلوكي", round(overall_stats.get("avg_behavioral") or 0, 1)])
    ws2.append(["متوسط المهاري", round(overall_stats.get("avg_skill") or 0, 1)])
    ws2.append(["متوسط الاجتماعي", round(overall_stats.get("avg_social") or 0, 1)])

    return _make_excel_response(wb, f"تقرير_جماعي_{date.today()}.xlsx")


def _export_group_csv(per_beneficiary, date_from, date_to):
    response = HttpResponse(content_type="text/csv; charset=utf-8-sig")
    response["Content-Disposition"] = f'attachment; filename="group_report_{date.today()}.csv"'
    response.write("﻿")
    writer = csv.writer(response)
    writer.writerow(["الاسم", "الرقم المرجعي", "المتوسط", "عدد التقييمات", "الطبي", "السلوكي", "المهاري", "الاجتماعي"])
    for p in per_beneficiary:
        writer.writerow([
            p["assessment__beneficiary__full_name"], p["assessment__beneficiary__ref_number"],
            round(p["avg_score"] or 0, 1), p["sessions"],
            round(p["avg_medical"] or 0, 1), round(p["avg_behavioral"] or 0, 1),
            round(p["avg_skill"] or 0, 1), round(p["avg_social"] or 0, 1),
        ])
    return response


def _export_period_excel(daily_stats, center_performance, overall, date_from, date_to):
    wb = Workbook()
    ws = wb.active
    ws.title = "الأداء اليومي"
    ws.sheet_view.rightToLeft = True
    ws.append(["التاريخ", "المتوسط", "عدد التقييمات", "تقدّم", "ثابت", "تراجع"])
    _style_header(ws)
    for d in daily_stats:
        ws.append([
            str(d["assessment__date"]),
            round(d["avg_score"] or 0, 1), d["count"],
            d["progressing"], d["stable"], d["regressing"],
        ])

    ws2 = wb.create_sheet("أداء المراكز")
    ws2.sheet_view.rightToLeft = True
    ws2.append(["المركز", "المحور", "متوسط الدرجة"])
    _style_header(ws2)
    for c in center_performance:
        ws2.append([c["name"], c["axis"], c["avg_score"]])

    ws3 = wb.create_sheet("ملخص")
    ws3.sheet_view.rightToLeft = True
    ws3.append(["الفترة", f"{date_from} إلى {date_to}"])
    ws3.append(["إجمالي التقييمات", overall.get("total", 0)])
    ws3.append(["المتوسط العام", round(overall.get("avg_total") or 0, 1)])

    return _make_excel_response(wb, f"تقرير_فترة_{date_from}_{date_to}.xlsx")


@require_perm("beneficiaries", "delete")
def beneficiary_data_delete(request, pk):
    b = get_object_or_404(Beneficiary, pk=pk)
    if request.method == "POST":
        DailyAssessment.objects.filter(beneficiary=b).delete()
        b.full_name = "محذوف"
        b.national_id = ""
        b.phone = ""
        b.emergency_contact = ""
        b.emergency_phone = ""
        b.address = ""
        b.admission_reason = ""
        b.notes = ""
        b.status = "discharged"
        if b.photo:
            b.photo.delete(save=False)
        b.save()
        from django.contrib import messages
        messages.success(request, "تم حذف بيانات المستفيد الشخصية بنجاح (حق الحذف)")
        return render(request, "reports/data_deleted.html", {"b": b})
    return render(request, "reports/data_delete_confirm.html", {"b": b})
